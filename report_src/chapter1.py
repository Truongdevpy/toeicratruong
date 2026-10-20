# -*- coding: utf-8 -*-
from docx.shared import Inches, Pt, RGBColor
from report_src.core import (
    add_heading_1, add_heading_2, add_body_p, add_bullet, 
    add_styled_table, add_table_caption, add_callout
)

def add_chapter_1(doc):
    add_heading_1(doc, "CHƯƠNG 1: GIỚI THIỆU")
    
    add_body_p(
        doc,
        "Trong xu thế hội nhập quốc tế và toàn cầu hóa mạnh mẽ hiện nay, năng lực giao tiếp tiếng Anh đã trở thành "
        "tiêu chuẩn bắt buộc đối với sinh viên tốt nghiệp các trường đại học, cao đẳng cũng như lực lượng lao động tri thức. "
        "Bài thi TOEIC (Test of English for International Communication) do Viện Khảo thí Giáo dục Hoa Kỳ (ETS) thiết kế "
        "hiện là chuẩn mực đánh giá năng lực tiếng Anh trong môi trường làm việc quốc tế được hơn 14.000 tổ chức, doanh nghiệp "
        "tại 160 quốc gia công nhận. Tại Việt Nam, hầu hết các trường đại học đều áp dụng chuẩn đầu ra TOEIC từ 450 đến 650+ điểm. "
        "Tuy nhiên, thực tế quá trình ôn luyện thi của sinh viên đang đối mặt với nhiều rào cản lớn về chi phí, phương pháp học "
        "và thiếu sự phản hồi thông minh. Chương này sẽ trình bày tổng quan bài toán, phân tích tình hình nghiên cứu hiện nay "
        "và xác định rõ mục tiêu, phạm vi nghiên cứu của đề tài 'Hệ thống luyện thi TOEIC trực tuyến chuẩn ETS tích hợp Trợ lý AI "
        "Socratic và Học tập thích ứng (TOEIC Master Platform)'."
    )

    add_heading_2(doc, "1.1 Tổng quan bài toán")
    add_body_p(
        doc,
        "1.1.1 Bối cảnh thực tiễn & Nhu cầu cấp bách của người học:",
        bold_prefix="a. Thực trạng đào tạo và chuẩn đầu ra: "
    )
    add_body_p(
        doc,
        "Theo thống kê khảo thí tại nhiều trường đại học khối kỹ thuật và kinh tế tại Việt Nam, tỷ lệ sinh viên năm cuối "
        "chưa thể nhận bằng tốt nghiệp đúng hạn do nợ chứng chỉ tiếng Anh chuẩn đầu ra (chủ yếu là TOEIC 450 - 650) dao động "
        "từ 20% đến 35%. Áp lực thời gian tốt nghiệp, yêu cầu tuyển dụng khắt khe từ các doanh nghiệp đa quốc gia khiến nhu cầu "
        "luyện thi TOEIC cấp tốc ngày càng tăng vọt."
    )
    add_body_p(
        doc,
        "1.1.2 Những rào cản và thách thức lớn trong phương pháp học truyền thống:",
        bold_prefix="b. Khó khăn của sinh viên khi tự ôn luyện: "
    )
    add_bullet(
        doc,
        "Học phí tại các trung tâm ngoại ngữ danh tiếng dao động từ 5.000.000đ đến 15.000.000đ cho một khóa ôn luyện từ mất gốc "
        "lên 650+, đây là khoản chi phí quá lớn đối với đại đa số sinh viên.",
        bold_prefix="• Gánh nặng chi phí kinh tế: "
    )
    add_bullet(
        doc,
        "Đề thi TOEIC Listening & Reading của ETS chứa vô số bẫy đánh lừa tinh vi: bẫy từ đồng âm (homophones), bẫy cùng gốc từ "
        "nhưng sai từ loại (word form), bẫy phủ định ngầm (subtle negatives), bẫy đảo ngữ và bẫy phân tán thông tin trong đoạn văn kép Part 7. "
        "Sinh viên khi tự giải đề trên sách giấy hoặc web tĩnh chỉ biết đáp án đúng là A/B/C/D mà không hiểu vì sao mình mắc bẫy.",
        bold_prefix="• Thiếu người chỉ điểm bẫy đề thi (Trap Analysis): "
    )
    add_bullet(
        doc,
        "Các giáo trình ôn thi truyền thống thiết kế một lộ trình cố định cho mọi đối tượng. Một học sinh yếu kỹ năng nghe hội thoại "
        "nhanh Part 3, Part 4 vẫn phải học dàn trải ngữ pháp Part 5 giống như học sinh yếu từ vựng thương mại. Sự thiếu vắng thuật toán "
        "thích ứng (Adaptive Learning) khiến người học lãng phí thời gian và dễ nản chí.",
        bold_prefix="• Lộ trình cào bằng, không cá nhân hóa: "
    )
    add_bullet(
        doc,
        "Tự học đòi hỏi tính tự giác cao. Không có hệ thống nhắc nhở, theo dõi chuỗi ngày học liên tục (streak) hay cơ chế khen thưởng "
        "tức thì (gamification), tỷ lệ bỏ cuộc giữa chừng của học viên tự học lên tới trên 70%.",
        bold_prefix="• Thiếu kỷ luật và động lực duy trì: "
    )

    add_heading_2(doc, "1.2 Tình hình nghiên cứu hiện tại")
    add_body_p(
        doc,
        "Nhằm đáp ứng nhu cầu tự học tiếng Anh trực tuyến, nhiều nền tảng công nghệ giáo dục (EdTech) đã ra đời trong những năm gần đây. "
        "Có thể phân loại các giải pháp hiện hành trên thị trường thành 3 nhóm chính: (1) Các website giải đề trắc nghiệm tĩnh; "
        "(2) Các nền tảng học trực tuyến quy mô lớn có trả phí; (3) Các ứng dụng luyện thi di động. Qua quá trình khảo sát thực tế chuyên sâu, "
        "nhóm nghiên cứu đã tổng hợp bảng so sánh tính năng toàn diện giữa giải pháp đề xuất (TOEIC Master Platform) và các hệ thống tiêu biểu:"
    )

    add_table_caption(doc, "1.1", "So sánh chi tiết giải pháp đề xuất với các nền tảng luyện thi hiện có")
    headers_c1 = [
        "Tiêu chí tính năng", 
        "Website trắc nghiệm tĩnh\n(ThiThuTOEIC, Violet)", 
        "Nền tảng thương mại lớn\n(Study4, Prep, DauTOEIC)", 
        "Hệ thống Đề xuất\n(TOEIC Master Platform)"
    ]
    data_c1 = [
        [
            "Mô phỏng phòng thi Full Test ETS",
            "Đồng hồ đếm ngược đơn giản, hay lỗi tải audio",
            "Mô phỏng khá tốt, có lưu lịch sử làm bài",
            "Chuẩn ETS 120 phút, audio đồng bộ câu, tua mượt mà HTTP Range"
        ],
        [
            "Trợ lý AI Socratic tương tác",
            "Hoàn toàn không có",
            "AI trả lời theo kịch bản cứng nhắc, mớm đáp án trực tiếp",
            "AI Socratic gợi mở tư duy, không đưa đáp án ngay, phản hồi theo bối cảnh câu hỏi"
        ],
        [
            "Phân tích bẫy đề thi (ETS Trap Analysis)",
            "Không hỗ trợ",
            "Chỉ có lời dịch văn bản hoặc bình luận của người dùng",
            "Tự động phát hiện bẫy (từ loại, đồng âm, bẫy thì) & sinh 'Mẹo 5 giây Socratic'"
        ],
        [
            "Lộ trình học thích ứng (Adaptive Learning)",
            "Không có",
            "Lộ trình cố định theo gói khóa học mua sẵn",
            "Động cơ thích ứng tự động tính trọng số (%) điểm yếu sau mỗi bài làm"
        ],
        [
            "Hạ tầng AI chịu tải cao (AI Pool)",
            "Không có",
            "Dễ quá tải hoặc giới hạn lượt hỏi khắt khe",
            "Gemini Key Pool xoay vòng tự động, hàng đợi Cooldown 65s chống lỗi 429"
        ],
        [
            "Luyện nghe chép chính tả (Dictation)",
            "Không có",
            "Có bài tập nghe nhưng ít đề cập nhật",
            "Chép chính tả thời gian thực, phát hiện sai sót từng ký tự & tính % chuẩn xác"
        ],
        [
            "Ghi nhớ từ vựng ngắt quãng (SRS)",
            "Không có",
            "Có Flashcard cơ bản",
            "Flashcard SRS khoa học (lặp lại ngắt quãng), phân nhóm từ vựng theo chủ đề ETS"
        ],
        [
            "Chi phí & Khả năng tiếp cận",
            "Miễn phí nhưng chất lượng thấp, nhiều quảng cáo",
            "Chi phí cao (800.000đ - 2.500.000đ/gói)",
            "Tiếp cận rộng rãi, mã nguồn mở, tối ưu cho sinh viên đại học"
        ]
    ]
    add_styled_table(doc, headers_c1, data_c1, col_widths=[1.5, 1.6, 1.8, 2.1])

    add_callout(
        doc,
        "Kết luận từ khảo sát: Các nền tảng hiện nay dù có nguồn đề phong phú nhưng đa phần vẫn hoạt động theo cơ chế thụ động "
        "(Passive Learning) – người học làm bài và đọc đáp án sẵn có. Sự đột phá của hệ thống TOEIC Master Platform nằm ở việc "
        "kết hợp Trí tuệ nhân tạo đàm thoại Socratic (chủ động gợi mở kiến thức), Bộ máy phân tích bẫy khảo thí độc quyền và "
        "Lộ trình thích ứng linh hoạt theo thời gian thực.",
        title="Nhận định cốt lõi:"
    )

    add_heading_2(doc, "1.3 Mục tiêu và phạm vi đề tài")
    add_body_p(
        doc,
        "Đề tài hướng tới việc giải quyết triệt để các rào cản học tập nêu trên thông qua các mục tiêu cụ thể sau:",
        bold_prefix="Mục tiêu nghiên cứu và phát triển: "
    )
    add_bullet(
        doc,
        "Thiết kế và lập trình hoàn chỉnh một Single Page Application (SPA) trên nền React 19 và Tailwind CSS v4, "
        "tối ưu hóa giao diện phòng thi chuẩn ETS, mượt mà trên cả trình duyệt máy tính và thiết bị di động.",
        bold_prefix="1. Phát triển nền tảng Full-stack hiện đại: "
    )
    add_bullet(
        doc,
        "Xây dựng dịch vụ quản lý cụm API Key của Google Gemini (`gemini_pool.py`) với cơ chế giám sát ngưỡng tải, "
        "tự động kích hoạt hàng đợi làm nguội khi gặp HTTP 429 Rate Limit và luân chuyển mô hình linh hoạt (Gemini 3 Flash, Gemini 2.5 Flash), "
        "đảm bảo hệ thống AI luôn phản hồi ổn định 24/7.",
        bold_prefix="2. Giải quyết bài toán chịu tải AI (Gemini Pool Engine): "
    )
    add_bullet(
        doc,
        "Triển khai thuật toán chấm thi TOEIC Listening & Reading chuẩn xác theo thang 10 - 990 điểm; "
        "tích hợp hệ thống phân tích lỗi sai và bẫy đề thi theo phương pháp Socratic Rule (quy tắc 5 giây).",
        bold_prefix="3. Khảo thí chuẩn ETS & Phân tích bẫy thông minh: "
    )
    add_bullet(
        doc,
        "Phát triển bài kiểm tra phân loại trình độ đầu vào (Placement Test) và động cơ thích ứng (Personalization Engine) "
        "liên tục tái tính toán tỷ trọng phân bổ bài học theo mức độ thành thạo của học viên.",
        bold_prefix="4. Lộ trình học thích ứng cá nhân hóa: "
    )
    add_bullet(
        doc,
        "Tích hợp đầy đủ các module: Luyện nghe chép chính tả (Dictation) sửa lỗi từng ký tự; Hệ thống thẻ ghi nhớ lặp lại ngắt quãng (SRS Flashcard); "
        "Luyện đề chuyên sâu từng phần (Part 1 - Part 7); Công cụ dịch nhanh khi bôi đen văn bản (Selection Translator).",
        bold_prefix="5. Hệ sinh thái bổ trợ toàn diện: "
    )
    add_body_p(
        doc,
        "Về phạm vi đề tài: Hệ thống tập trung chuyên sâu vào bài thi TOEIC 2 kỹ năng cốt lõi (Listening & Reading) "
        "với hơn 10 bộ đề Full Test hoàn chỉnh (200 câu/đề) kèm hàng ngàn câu hỏi phân loại theo Part 1 - Part 7. "
        "Đồng thời, hệ thống cung cấp sẵn khung sườn và đề thi mẫu cho 2 kỹ năng Speaking & Writing để học viên rèn luyện toàn diện.",
        bold_prefix="Phạm vi ứng dụng: "
    )

    add_heading_2(doc, "1.4 Đối tượng thụ hưởng")
    add_body_p(
        doc,
        "Sản phẩm của đề tài được thiết kế nhằm phục vụ trực tiếp các nhóm đối tượng cụ thể:",
        bold_prefix="Phân khúc người dùng mục tiêu: "
    )
    add_bullet(
        doc,
        "Sinh viên các trường đại học, cao đẳng trên toàn quốc đang chuẩn bị thi chứng chỉ chuẩn đầu ra tốt nghiệp, "
        "đặc biệt là sinh viên khối kỹ thuật, sư phạm, kinh tế tự học tại nhà cần lộ trình chuẩn và chi phí tiết kiệm.",
        bold_prefix="• Sinh viên năm cuối đại học: "
    )
    add_bullet(
        doc,
        "Những người đi làm bận rộn cần nâng cao band điểm TOEIC lên 750+ để thăng tiến hoặc ứng tuyển vào các tập đoàn đa quốc gia, "
        "có nhu cầu học linh hoạt trong các khoảng thời gian trống 15-30 phút mỗi ngày.",
        bold_prefix="• Người đi làm (Working Professionals): "
    )
    add_bullet(
        doc,
        "Giáo viên có thể sử dụng hệ thống như một ngân hàng đề thi chuẩn hóa, công cụ giao bài tập và theo dõi biểu đồ tiến độ "
        "của học sinh mà không phải tốn hàng giờ chấm bài thủ công.",
        bold_prefix="• Giáo viên và Giảng viên tiếng Anh: "
    )
    doc.add_page_break()
