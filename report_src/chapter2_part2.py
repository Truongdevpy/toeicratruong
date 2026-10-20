# -*- coding: utf-8 -*-
from docx.shared import Inches, Pt, RGBColor
from report_src.core import (
    add_heading_3, add_body_p, 
    add_bullet, add_styled_table, add_table_caption, add_figure_caption
)

def add_chapter_2_part2(doc):
    # 2.2.3 Lộ trình thích ứng
    add_heading_3(doc, "2.2.3 Động cơ Luyện tập thích ứng & Lộ trình cá nhân hóa")
    add_body_p(
        doc,
        "Động cơ thích ứng (personalization_engine.py, routers/learning_path.py) là trái tim điều phối học tập của hệ thống. "
        "Thay vì một lộ trình học tuyến tính cố định, hệ thống vận hành theo thuật toán thích ứng liên tục:",
        bold_prefix="Cơ chế học tập thích ứng (Adaptive Learning): "
    )
    add_bullet(
        doc,
        "Học viên mới được cung cấp bài kiểm tra phân loại nhanh gồm các câu hỏi đại diện cho 7 Part của đề thi TOEIC. "
        "Dựa trên kết quả bài làm, hệ thống tự động suy diễn trình độ theo khung tham chiếu Châu Âu CEFR (từ A1 đến B2/C1) "
        "và ước lượng band điểm xuất phát (ví dụ: 425 điểm).",
        bold_prefix="• Kiểm tra đánh giá đầu vào (Placement Test): "
    )
    add_bullet(
        doc,
        "Hệ thống định nghĩa một ma trận trọng số phân bổ thời gian học tập ban đầu: Grammar 35%, Listening 30%, Reading 20%, Vocabulary 15%. "
        "Sau mỗi bài kiểm tra hoặc bài luyện tập, thuật toán phân tích tỷ lệ làm đúng trên từng Part. Nếu một kỹ năng (ví dụ: Word Form Part 5 hoặc Đọc hiểu Part 7) "
        "có tỷ lệ đúng dưới 45%, thuật toán sẽ tự động điều chỉnh tăng trọng số phân bổ cho kỹ năng đó, đồng thời mở khóa các Nút học tập (Learning Nodes) "
        "chuyên đề tương ứng trong bảng learning_nodes.",
        bold_prefix="• Thuật toán tính trọng số thích ứng (Adaptive Weight Calculation): "
    )

    add_table_caption(doc, "2.3", "Bảng đặc tả Use Case Chức năng Đánh giá đầu vào & Sinh lộ trình thích ứng")
    headers_uc3 = ["Thuộc tính Use Case", "Mô tả chi tiết kỹ thuật"]
    data_uc3 = [
        ["Tên Use Case", "Đánh giá đầu vào & Sinh lộ trình học thích ứng (placement.py & learning_path.py)"],
        ["Tác nhân", "Học viên (Student), Động cơ thích ứng (System)"],
        ["Mô tả", "Kiểm tra trình độ xuất phát, xác định mục tiêu và tự động sinh cây lộ trình học thích ứng cá nhân hóa."],
        ["Tiền điều kiện", "Học viên hoàn tất đăng ký tài khoản."],
        ["Luồng sự kiện chính", 
         "1. Học viên làm bài kiểm tra Placement Test.\n"
         "2. Hệ thống chấm điểm, đối chiếu khung CEFR và tính điểm TOEIC tương đương.\n"
         "3. Hệ thống lưu hồ sơ học viên vào bảng user_profiles.\n"
         "4. Thuật toán thích ứng khởi tạo bản ghi learning_paths và sinh chuỗi learning_nodes.\n"
         "5. Mỗi nút kỹ năng được gắn trạng thái ('locked', 'in_progress', 'completed') và % tiến độ.\n"
         "6. Khi học viên hoàn thành bài tập, điểm yếu mới được cập nhật và lộ trình tự động tái cân bằng."],
        ["Hậu điều kiện", "Cây lộ trình cá nhân hóa được hiển thị trực quan trên Dashboard của học viên."]
    ]
    add_styled_table(doc, headers_uc3, data_uc3, col_widths=[2.0, 4.5])
    add_figure_caption(doc, "2.6", "Sơ đồ thuật toán tính trọng số thích ứng (Adaptive Learning Engine)")

    # 2.2.4 Dictation & Flashcard
    add_heading_3(doc, "2.2.4 Luyện nghe chép chính tả (Dictation) & Flashcard lặp lại ngắt quãng (SRS)")
    add_body_p(
        doc,
        "Nhằm củng cố khả năng phản xạ nghe và mở rộng vốn từ vựng học thuật, hệ thống tích hợp hai công cụ bổ trợ đắc lực:",
        bold_prefix="Công cụ gia tăng phản xạ ngôn ngữ: "
    )
    add_bullet(
        doc,
        "Được xây dựng trong routers/dictation.py. Học viên nghe từng câu hội thoại thực tế được cắt gọt chuẩn xác từ đề thi ETS, "
        "sau đó gõ lại chuỗi văn bản nghe được vào ô nhập liệu. Hệ thống áp dụng thuật toán so khớp chuỗi ký tự theo thời gian thực "
        "(Character-level String Matching), tự động làm nổi bật các từ bị bỏ sót, từ viết sai chính tả và tính toán tỷ lệ phần trăm chính xác (Accuracy %). "
        "Phương pháp này giúp học viên khắc phục triệt để hiện tượng nuốt âm, nối âm trong đề thi TOEIC Listening.",
        bold_prefix="• Module Luyện nghe chép chính tả (Dictation Engine): "
    )
    add_bullet(
        doc,
        "Được triển khai trong routers/vocabulary.py dựa trên nguyên lý Đường cong quên lãng của Ebbinghaus và thuật toán SuperMemo (SM-2). "
        "Mỗi từ vựng mới được lưu trong bảng user_flashcards với các thuộc tính: số lần ôn tập (repetition_count), khoảng cách ngày lặp lại (interval_days) "
        "và hệ số ghi nhớ (ease_factor). Các từ học viên đánh giá 'Khó' sẽ xuất hiện lặp lại thường xuyên hơn, giúp từ vựng được khắc sâu vào trí nhớ dài hạn.",
        bold_prefix="• Thẻ ghi nhớ lặp lại ngắt quãng (Spaced Repetition System - SRS Flashcards): "
    )

    # 2.2.5 Hệ thống Crawler & Chuẩn hóa dữ liệu
    add_heading_3(doc, "2.2.5 Hệ thống thu thập và chuẩn hóa dữ liệu đề thi quy mô lớn")
    add_body_p(
        doc,
        "Để đảm bảo nguồn tài liệu phong phú và cập nhật nhất, hệ sinh thái sở hữu phân hệ cào và bóc tách dữ liệu tự động (data/crawler.py):",
        bold_prefix="Quy trình thu thập và chuẩn hóa dữ liệu đề thi: "
    )
    add_bullet(
        doc,
        "Tự động thu thập đề thi từ các kho tài liệu uy tín (Đấu TOEIC, Study4). Bóc tách cấu trúc câu hỏi, câu trả lời, lời thoại transcript, "
        "đáp án chính xác và lời giải thích tiếng Việt chi tiết.",
        bold_prefix="• Tự động hóa bóc tách cấu trúc đề thi: "
    )
    add_bullet(
        doc,
        "Mọi tệp âm thanh MP3 và hình ảnh đề thi tải về đều được tính toán mã băm SHA-256 lưu trong cơ sở dữ liệu crawler.db. "
        "Nếu một đoạn âm thanh hướng dẫn chung hay hình ảnh xuất hiện lặp lại ở nhiều bài thi, hệ thống tự động tái sử dụng đường dẫn lưu trữ, "
        "giúp tiết kiệm hơn 40% dung lượng ổ đĩa và ngăn chặn hoàn toàn việc tải trùng lặp.",
        bold_prefix="• Cơ chế chống trùng lặp tệp bằng mã băm SHA-256 (Deduplication): "
    )
    add_bullet(
        doc,
        "Toàn bộ dữ liệu sau khi làm sạch được biên tập thành định dạng Web-Ready JSON (tests_catalog.json, practice_by_part/part1..7.json), "
        "giúp máy chủ API nạp dữ liệu vào bộ nhớ đệm siêu tốc mà không cần thực hiện các truy vấn lồng phức tạp.",
        bold_prefix="• Xuất dữ liệu tối ưu hóa cho Web (Web-Ready Datasets): "
    )

    # 2.2.6 An toàn thông tin
    add_heading_3(doc, "2.2.6 Cơ chế An toàn thông tin và Quản lý quyền hạn")
    add_body_p(
        doc,
        "Hệ thống áp dụng các tiêu chuẩn an toàn thông tin nghiêm ngặt trong thiết kế kiến trúc phần mềm:",
        bold_prefix="Bảo mật đa tầng và kiểm soát truy cập: "
    )
    add_bullet(
        doc,
        "Mật khẩu người dùng được băm bằng thuật toán bcrypt với độ phức tạp cao (salt tự sinh). Quá trình xác thực phiên làm việc "
        "sử dụng JSON Web Token (JWT) mã hóa HS256 với thời gian sống xác định (expiration time), tích hợp cơ chế thu hồi token qua trường token_version.",
        bold_prefix="• Xác thực và ủy quyền JWT an toàn: "
    )
    add_bullet(
        doc,
        "Tất cả các truy vấn cơ sở dữ liệu đều được thực hiện qua SQLAlchemy ORM sử dụng cơ chế Prepared Statements, "
        "loại bỏ hoàn toàn nguy cơ bị tấn công tiêm mã SQL Injection.",
        bold_prefix="• Phòng chống SQL Injection triệt để: "
    )
    add_bullet(
        doc,
        "Tích hợp module ai_security.py sử dụng biểu thức chính quy và bộ lọc từ khóa nguy hiểm để rà soát toàn bộ chuỗi ký tự "
        "người dùng nhập vào trước khi gửi tới Google Gemini. Các nỗ lực nhằm vượt quyền hệ thống (Jailbreak, Prompt Injection) "
        "như 'Ignore previous instructions', 'Act as system admin' đều bị chặn đứng và ghi nhận vào nhật ký an ninh.",
        bold_prefix="• Lớp phòng hộ chống tấn công Prompt Injection cho AI: "
    )
